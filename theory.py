'''
 Code written by Sarthak Mishra. 
 Refactored and edited for readability with assistance from OpenAI Codex. 
'''

import argparse
import csv
import math
from pathlib import Path


def int_exp(rate, upper):
    """Integral_0^upper exp(-rate*t) dt."""
    if abs(rate) < 1e-12:
        return upper
    return -math.expm1(-rate * upper) / rate


def int_t_exp(rate, upper):
    """Integral_0^upper t*exp(-rate*t) dt."""
    if abs(rate) < 1e-12:
        return upper * upper / 2.0
    x = rate * upper
    return (1.0 - math.exp(-x) * (1.0 + x)) / (rate * rate)


def evaluate_branch_lengths(ts, lambda_rate):
    if ts <= 0:
        raise ValueError("ts must be > 0.")
    if lambda_rate <= 0:
        raise ValueError("lambda must be > 0 because the section conditions on one duplication.")
    if abs(lambda_rate - 1.0) < 1e-10:
        raise ValueError(
            "lambda = 1 is a removable singular case in several manuscript expressions. "
            "Use a value slightly different from 1, or extend the script with the lambda->1 limits."
        )

    lam = lambda_rate

    # Tree 1: ((A1,B1),(A2,B2)) -- manuscript regime 1
    tree1_b1 = 1.0
    tree1_b2 = 1.0 + 1.0 / (1.0 + 3.0 * lam) + 1.0 / (3.0 + 2.0 * lam)

    # Tree 2-3: ((A1,(A2,B2)),B1)
    tree2_lower = 1.0 / (1.0 + 3.0 * lam) + 1.0 / (3.0 + 2.0 * lam)
    tree2_upper = 1.0 / (1.0 + 2.0 * lam)

    # Tree 4-5: ((B1,B2),A1)
    i0_lam = int_exp(lam, ts)
    i0_1 = int_exp(1.0, ts)
    i1_lam = int_t_exp(lam, ts)
    i1_1 = int_t_exp(1.0, ts)

    denom_tk = i0_lam - i0_1
    if abs(denom_tk) < 1e-14:
        raise ValueError("Regime-1 E[t_k] denominator is numerically zero for these inputs.")

    etk_r1 = (i1_lam - i1_1) / denom_tk
    tree3_r1 = ts - etk_r1 + 1.0 / (1.0 + 2.0 * lam)
    tree3_r2 = 1.0 / (1.0 + 2.0 * lam)
    tree3_r3 = 1.0 / (1.0 + 2.0 * lam)

    # Consistent weights from the earlier regime posterior kernels.
    a = 1.0 + 2.0 * lam
    c = 1.0 - lam
    q = 3.0 + 2.0 * lam
    r = 2.0 + 3.0 * lam

    p1 = math.exp(-2.0 * lam * ts) / (a * c) * (i0_lam - i0_1)
    f_ts = math.expm1(c * ts) / c
    p2 = math.exp(-(1.0 + 2.0 * lam) * ts) / (a * q) * f_ts
    p3 = math.exp(-3.0 * lam * ts) / (a * c) * (1.0 / r - 1.0 / q)

    total_p = p1 + p2 + p3
    w1, w2, w3 = p1 / total_p, p2 / total_p, p3 / total_p
    tree3_overall = tree3_r1 * w1 + tree3_r2 * w2 + tree3_r3 * w3

    # Tree 6-7: ((A1,B1),B2)
    tree4_b = 1.0

    # Tree 8-9: ((A1,B2),B1)
    tree5_b = 1.0 / (1.0 + 2.0 * lam)

    rows = [
        ["((A1,B1),(A2,B2))", "b1", "Only regime", tree1_b1,
         "Equation 24"],
        ["((A1,B1),(A2,B2))", "b2", "Only regime", tree1_b2,
         "Equation 27"],
        ["((A1,(A2,B2)),B1)", "b_lower", "Only regime", tree2_lower,
         "Equation 42"],
        ["((A1,(A2,B2)),B1)", "b_upper", "Only regime", tree2_upper,
         "Equation 44"],
        ["((B1,B2),A1)", "b", "Regime 1", tree3_r1,
         "Equation 60"],
        ["((B1,B2),A1)", "b", "Regime 2", tree3_r2,
         "Equation 64"],
        ["((B1,B2),A1)", "b", "Regime 3", tree3_r3,
         "Equation 68"],
        ["((B1,B2),A1)", "b", "Weighted overall", tree3_overall,
         "Equation 70"],
        ["((A1,B1),B2)", "b", "Only regime", tree4_b,
         "Equation 84"],
        ["((A1,B2),B1)", "b", "Overall", tree5_b,
         "Equation 100"],
    ]

    weights = [
        ["Regime 1", p1, w1],
        ["Regime 2", p2, w2],
        ["Regime 3", p3, w3],
    ]

    return rows, weights


RESULT_HEADERS = ["Tree", "Branch", "Regime", "Expected branch length", "Notes"]
WEIGHT_HEADERS = ["Regime", "Unnormalized weight", "Normalized weight"]


def write_csv(ts, lambda_rate, output_path):
    """Write results and a companion *_weights.csv using the standard library."""
    rows, weights = evaluate_branch_lengths(ts, lambda_rate)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    weights_path = output_path.with_name(output_path.stem + "_weights.csv")
    for path, headers, data in (
        (output_path, RESULT_HEADERS, rows),
        (weights_path, WEIGHT_HEADERS, weights),
    ):
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(data)
    return weights_path


def write_excel(ts, lambda_rate, output_path):
    """Write an Excel workbook using the optional openpyxl package."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows, weights = evaluate_branch_lengths(ts, lambda_rate)
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    for row in [
        ["Branch Length Evaluator", ""],
        ["Parameter", "Value"],
        ["t_s", ts],
        ["lambda", lambda_rate],
    ]:
        inputs.append(row)
    inputs.merge_cells("A1:B1")
    results = wb.create_sheet("Results")
    results.append(RESULT_HEADERS)
    for row in rows:
        results.append(row)
    results.freeze_panes = "A2"
    calc = wb.create_sheet("Tree3_Weights")
    calc.append(WEIGHT_HEADERS)
    for row in weights:
        calc.append(row)

    for sheet in wb:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
        for column in sheet.iter_cols():
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 42)
    results.column_dimensions["A"].width = 24
    results.column_dimensions["E"].width = 42
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--ts", type=float, required=True, help="Speciation time t_s")
    parser.add_argument("--lambda-rate",dest="lambda_rate",type=float,nargs="+",default=[i / 100 for i in range(1, 11)],
        help="One or more duplication rates (default: 0.01 to 0.10 in steps of 0.01)")
    parser.add_argument("--output",default="branch_length_results.csv",
        help="Output .csv or .xlsx filename; multiple rates add a _lambda suffix. Use {lambda} to place the rate explicitly. .xlsx requires openpyxl",
    )
    args = parser.parse_args()

    suffix = Path(args.output).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        parser.error("Output filename must end in .csv or .xlsx.")

    for lambda_rate in args.lambda_rate:
        rate_label = str(lambda_rate)
        output_path = Path(args.output.replace("{lambda}", rate_label))
        if len(args.lambda_rate) > 1 and "{lambda}" not in args.output:
            output_path = output_path.with_name(
                f"{output_path.stem}_{rate_label}{output_path.suffix}"
            )
        if suffix == ".csv":
            weights_path = write_csv(args.ts, lambda_rate, output_path)
            print(f"Saved weights: {weights_path}")
        else:
            try:
                write_excel(args.ts, lambda_rate, output_path)
            except ModuleNotFoundError as exc:
                if exc.name != "openpyxl":
                    raise
                parser.error("Excel output requires openpyxl. Install it with 'python -m pip install openpyxl' or use --output results.csv.")
        print(f"Saved: {output_path}")



if __name__ == "__main__":
    main()
